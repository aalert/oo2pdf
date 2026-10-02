"""Generate DOCX/XLSX fixtures used for fidelity testing (dev only)."""
from __future__ import annotations

import datetime as dt
import io
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "fixtures"


def _png(w=320, h=180, color=(70, 130, 180)):
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (w, h), color)
    d = ImageDraw.Draw(im)
    d.ellipse((w * 0.3, h * 0.2, w * 0.7, h * 0.8), fill=(250, 200, 60))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    return buf


LOREM = ("Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor "
         "incididunt ut labore et dolore magna aliqua. Ut enim ad minim veniam, quis nostrud "
         "exercitation ullamco laboris nisi ut aliquip ex ea commodo consequat. Duis aute irure "
         "dolor in reprehenderit in voluptate velit esse cillum dolore eu fugiat nulla pariatur.")


def make_basic_docx():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.shared import Inches, Pt, RGBColor

    d = Document()
    d.add_heading("Quarterly Report", 0)
    d.add_heading("1. Introduction", 1)
    p = d.add_paragraph(LOREM)
    p = d.add_paragraph()
    p.add_run("Bold text, ").bold = True
    p.add_run("italic text, ").italic = True
    r = p.add_run("underlined text, ")
    r.underline = True
    r = p.add_run("red 14pt text")
    r.font.color.rgb = RGBColor(0xC0, 0, 0)
    r.font.size = Pt(14)
    p.add_run(" and E = mc")
    r = p.add_run("2")
    r.font.superscript = True
    p.add_run(". Then H")
    r = p.add_run("2")
    r.font.subscript = True
    p.add_run("O and some ")
    r = p.add_run("highlighted")
    r.font.highlight_color = 7  # yellow
    p.add_run(" words.")

    p = d.add_paragraph(LOREM + " " + LOREM)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p = d.add_paragraph("Centered paragraph with Times New Roman 12pt.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
    p = d.add_paragraph("Right aligned paragraph in Arial 10pt.")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in p.runs:
        r.font.name = "Arial"
        r.font.size = Pt(10)

    d.add_heading("2. Lists", 1)
    for t in ("First bullet item", "Second bullet item that is long enough to wrap onto a second "
              "line to check the hanging indent alignment of the list", "Third bullet"):
        d.add_paragraph(t, style="List Bullet")
    for t in ("Step one", "Step two", "Step three"):
        d.add_paragraph(t, style="List Number")

    d.add_heading("3. Tabs", 1)
    p = d.add_paragraph()
    ts = p.paragraph_format.tab_stops
    ts.add_tab_stop(Inches(3), WD_TAB_ALIGNMENT.CENTER)
    ts.add_tab_stop(Inches(6), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p.add_run("Left\tCenter\tRight")
    p = d.add_paragraph("Chapter One\t12")
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p = d.add_paragraph("Default\ttab\tstops\tevery\thalf\tinch")

    d.add_heading("4. Tables", 1)
    t = d.add_table(rows=4, cols=3)
    t.style = "Table Grid"
    data = [("Region", "Units", "Revenue"), ("North", "1,204", "$45,300"),
            ("South", "987", "$38,120"), ("West", "1,530", "$61,002")]
    for i, row in enumerate(data):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    d.add_paragraph()
    t = d.add_table(rows=5, cols=4)
    t.style = "Light Grid Accent 1"
    for i in range(5):
        for j in range(4):
            t.cell(i, j).text = f"R{i}C{j}" if i else f"Header {j}"
    t.cell(1, 0).merge(t.cell(2, 0))
    t.cell(3, 1).merge(t.cell(3, 2))
    d.add_paragraph()

    d.add_heading("5. Picture", 1)
    d.add_picture(_png(), width=Inches(3))
    d.add_paragraph("Caption: a generated picture.", style="Caption")

    p = d.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    d.add_heading("6. Second page", 1)
    for i in range(12):
        d.add_paragraph(f"Paragraph {i + 1}. " + LOREM)

    # header / footer with page numbers
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "ACME Corp — Confidential"
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Page ")
    _add_field(fp, "PAGE")
    fp.add_run(" of ")
    _add_field(fp, "NUMPAGES")
    d.save(OUT / "basic.docx")


def _add_field(paragraph, code):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    run._r.append(b)
    run = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {code} "
    run._r.append(it)
    run = paragraph.add_run()
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    run._r.append(s)
    paragraph.add_run("1")
    run = paragraph.add_run()
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    run._r.append(e)


def make_long_docx():
    """Many paragraphs of varied spacing to stress pagination."""
    from docx import Document
    from docx.shared import Pt
    d = Document()
    for i in range(1, 9):
        d.add_heading(f"Section {i}", 1)
        for j in range(i % 3 + 2):
            d.add_paragraph(f"{i}.{j} " + LOREM * (1 + (i + j) % 3))
        if i % 2 == 0:
            d.add_heading(f"Subsection {i}.1", 2)
            p = d.add_paragraph("Line spacing double. " + LOREM)
            p.paragraph_format.line_spacing = 2.0
            p = d.add_paragraph("Exactly 18pt. " + LOREM)
            p.paragraph_format.line_spacing = Pt(18)
            p = d.add_paragraph("Space before 24 after 6. " + LOREM)
            p.paragraph_format.space_before = Pt(24)
            p.paragraph_format.space_after = Pt(6)
    d.save(OUT / "long.docx")


def make_basic_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"] = "Sales Report 2024"
    ws["A1"].font = Font(size=16, bold=True, color="1F4E78")
    ws.merge_cells("A1:E1")
    ws["A1"].alignment = Alignment(horizontal="center")
    hdr = ["Date", "Product", "Units", "Price", "Total"]
    thin = Side(style="thin", color="000000")
    for i, h in enumerate(hdr, 1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4472C4")
        c.alignment = Alignment(horizontal="center")
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    products = ["Widget", "Gadget", "Doohickey", "Thingamajig", "Whatsit"]
    for r in range(4, 44):
        i = r - 4
        ws.cell(row=r, column=1, value=dt.date(2024, 1, 1) + dt.timedelta(days=i * 3)).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=2, value=products[i % 5])
        ws.cell(row=r, column=3, value=(i * 7) % 23 + 1)
        c = ws.cell(row=r, column=4, value=round(9.99 + i * 1.37, 2))
        c.number_format = '"$"#,##0.00'
        c = ws.cell(row=r, column=5, value=f"=C{r}*D{r}")
        c.number_format = '#,##0.00;[Red]-#,##0.00'
        for col in range(1, 6):
            ws.cell(row=r, column=col).border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws["G3"] = "This is a long text that overflows into the neighbouring empty cells"
    ws["G5"] = 0.1234
    ws["G5"].number_format = "0.0%"
    ws["G6"] = 1234567.891
    ws["G7"] = 1234567.891
    ws["G7"].number_format = "#,##0"
    ws["G8"] = -42.5
    ws["G8"].number_format = '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)'
    ws["G9"] = "Wrapped text in a cell that is narrow"
    ws["G9"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["G"].width = 18
    ws.row_dimensions[9].height = 45
    ws["G10"] = dt.datetime(2024, 3, 15, 14, 30)
    ws["G10"].number_format = "dddd, mmmm d, yyyy h:mm AM/PM"
    ws["G11"] = 3.14159
    ws["G11"].number_format = "# ?/?"
    ws["G12"] = 0.000012345
    ws["G12"].number_format = "0.00E+00"
    ws.oddFooter.center.text = "Page &P of &N"
    ws.oddHeader.left.text = "&F"
    ws.oddHeader.right.text = "&A"

    ws2 = wb.create_sheet("Wide")
    for r in range(1, 61):
        for c in range(1, 16):
            ws2.cell(row=r, column=c, value=r * c)
    ws2.print_title_rows = "1:1"
    ws2.page_setup.orientation = "landscape"
    wb.save(OUT / "basic.xlsx")
    # compute formula cached values through Excel when making references


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    make_basic_docx()
    make_long_docx()
    make_basic_xlsx()
    print("fixtures written to", OUT)
