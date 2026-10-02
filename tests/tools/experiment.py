"""Run a Word/Excel layout experiment (dev only).

Builds documents with a python function, exports them through the installed
Office (reference) and through oo2pdf, then prints line baselines side by
side.  Used to discover Office's layout rules empirically.

    python tests/tools/experiment.py pagebreak
"""
from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(Path(__file__).parent))

OUT = Path(tempfile.gettempdir()) / "oo2pdf-exp"


def word_export(paths):
    lines = ["$w = New-Object -ComObject Word.Application", "$w.DisplayAlerts = 0"]
    for p in paths:
        pdf = p.with_suffix(".ref.pdf")
        lines.append(f"$d = $w.Documents.Open('{p}', $false, $true, $false); "
                     f"$d.ExportAsFixedFormat('{pdf}', 17); $d.Close(0)")
    lines.append("$w.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", "; ".join(lines)], check=True, timeout=300)


def excel_export(paths):
    lines = ["$x = New-Object -ComObject Excel.Application", "$x.DisplayAlerts = $false"]
    for p in paths:
        pdf = p.with_suffix(".ref.pdf")
        lines.append(f"$b = $x.Workbooks.Open('{p}'); $b.Save(); $b.ExportAsFixedFormat(0, '{pdf}'); $b.Close($false)")
    lines.append("$x.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", "; ".join(lines)], check=True, timeout=300)


def show(ref_pdf, ours_pdf, pages=None):
    from lines import lines
    import pymupdf
    n = len(pymupdf.open(ref_pdf))
    m = len(pymupdf.open(ours_pdf))
    print(f"pages ref={n} ours={m}")
    for pno in range(max(n, m)):
        if pages and pno + 1 not in pages:
            continue
        print(f"--- page {pno + 1}")
        R = lines(ref_pdf, pno) if pno < n else []
        O = lines(ours_pdf, pno) if pno < m else []
        for i in range(max(len(R), len(O))):
            r = R[i] if i < len(R) else (0, 0, "")
            o = O[i] if i < len(O) else (0, 0, "")
            s = f"{r[0]:8.2f} {r[1]:7.2f} {r[2]:<34} | {o[0]:8.2f} {o[1]:7.2f} {o[2]:<34}  dy={o[0]-r[0]:+.2f}"
            print(s.encode("ascii", "replace").decode())


def run(name, build, kind="docx", pages=None):
    import oo2pdf
    OUT.mkdir(exist_ok=True)
    path = OUT / f"{name}.{kind}"
    build(path)
    (word_export if kind == "docx" else excel_export)([path])
    ours = path.with_suffix(".ours.pdf")
    oo2pdf.convert(path, ours, **({"locale": "system"} if kind == "xlsx" else {}))
    show(path.with_suffix(".ref.pdf"), ours, pages)
    return path


# ------------------------------------------------------------------ experiments
def exp_pagebreak(path):
    from docx import Document
    from docx.enum.text import WD_BREAK
    d = Document()
    d.add_paragraph("P1 text")
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.add_heading("H after break paragraph", 1)
    d.add_paragraph("body")
    p = d.add_paragraph("P2 text then break")
    p.add_run().add_break(WD_BREAK.PAGE)
    d.add_paragraph("N after inline break")
    p = d.add_paragraph("P3 text")
    r = p.add_run()
    r.add_break(WD_BREAK.PAGE)
    p.add_run("tail after break")
    d.add_paragraph("N next")
    h = d.add_heading("H page break before", 1)
    h.paragraph_format.page_break_before = True
    d.add_paragraph("body2")
    d.save(path)




def exp_xlsizes(path):
    from openpyxl import Workbook
    from openpyxl.styles import Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    thin = Side(style="thin")
    widths = [1, 2, 3, 5, 8.43, 10, 13, 20, 30]
    heights = [5, 10, 12.75, 15, 20, 30, 45, 60, 100, 150]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for j, h in enumerate(heights, 1):
        ws.row_dimensions[j].height = h
    for j in range(1, len(heights) + 1):
        for i in range(1, len(widths) + 1):
            ws.cell(row=j, column=i).border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A12"] = "x"
    ws.page_margins.left = ws.page_margins.right = 0.5
    wb.save(path)


def exp_xltext(path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Side, Font
    wb = Workbook()
    ws = wb.active
    thin = Side(style="thin")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 20
    rows = [
        (1, "0"), (2, "0_)"), (3, "0_W"), (4, '0" "'), (5, "_(0"), (6, "0.00_);(0.00)"),
        (7, "#,##0.00_);[Red](#,##0.00)"), (-8, "#,##0.00_);[Red](#,##0.00)"), (9, "General"),
    ]
    for i, (v, f) in enumerate(rows, 1):
        c = ws.cell(row=i, column=1, value=v)
        c.number_format = f
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        ws.cell(row=i, column=2, value=f)
    for i, (h, v) in enumerate([("left", "Left"), ("center", "Center"), ("right", "Right"), (None, "General")], 12):
        c = ws.cell(row=i, column=1, value=v)
        c.alignment = Alignment(horizontal=h)
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, v in enumerate(["top", "center", "bottom"], 17):
        ws.row_dimensions[i].height = 40
        c = ws.cell(row=i, column=1, value=v)
        c.alignment = Alignment(vertical=v)
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    c = ws.cell(row=21, column=1, value="Big 20pt")
    c.font = Font(size=20)
    c = ws.cell(row=22, column=1, value="Indent 2")
    c.alignment = Alignment(indent=2)
    c = ws.cell(row=23, column=1, value="Wrapped text that needs several lines to fit in the cell")
    c.alignment = Alignment(wrap_text=True)
    c = ws.cell(row=24, column=1, value=1.234e-05)
    c = ws.cell(row=25, column=1, value=-0.0001)
    c.number_format = "0.00"
    c = ws.cell(row=26, column=1, value=0.999999)
    c.number_format = "hh:mm:ss"
    wb.save(path)


def exp_imgline(path):
    import io
    from docx import Document
    from docx.shared import Pt, Inches
    from PIL import Image
    im = Image.new("RGB", (200, 100), (200, 50, 50))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    d = Document()
    for label, ls in [("single", 1.0), ("m115", 1.15), ("double", 2.0), ("exact", Pt(14)), ("least", None)]:
        p = d.add_paragraph(f"A {label}")
        p.paragraph_format.space_after = Pt(0)
        buf.seek(0)
        p = d.add_paragraph()
        p.add_run().add_picture(buf, width=Inches(1))
        p.paragraph_format.space_after = Pt(0)
        if ls is None:
            p.paragraph_format.line_spacing_rule = 3  # at least
            p.paragraph_format.line_spacing = Pt(12)
        else:
            p.paragraph_format.line_spacing = ls
        p = d.add_paragraph(f"B {label}")
        p.paragraph_format.space_after = Pt(0)
    # picture with text on the same line
    p = d.add_paragraph("Text before ")
    buf.seek(0)
    p.add_run().add_picture(buf, width=Inches(0.5))
    p.add_run(" text after")
    d.add_paragraph("C after mixed")
    d.save(path)


def exp_rowborders(path):
    """Row heights with borders between rows: which side reserves the width?"""
    from docx import Document
    from docx.oxml import parse_xml
    from docx.shared import Pt
    d = Document()
    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    b12 = '<w:{side} w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
    cases = [("none", "", ""), ("r1bottom", b12.format(side="bottom"), ""),
             ("r2top", "", b12.format(side="top")),
             ("both", b12.format(side="bottom"), b12.format(side="top"))]
    for name, r1, r2 in cases:
        p = d.add_paragraph(f"before {name}")
        p.paragraph_format.space_after = Pt(0)
        rows = []
        for i, bd in enumerate((r1, r2)):
            rows.append(f'<w:tr><w:tc><w:tcPr><w:tcW w:w="4000" w:type="dxa"/><w:tcBorders>{bd}</w:tcBorders></w:tcPr>'
                        f'<w:p><w:pPr><w:spacing w:after="0" w:line="280" w:lineRule="exact"/></w:pPr>'
                        f'<w:r><w:t>{name} row{i + 1}</w:t></w:r></w:p></w:tc></w:tr>')
        tbl = parse_xml(f'<w:tbl {W}><w:tblPr><w:tblW w:w="0" w:type="auto"/></w:tblPr>'
                        f'<w:tblGrid><w:gridCol w:w="4000"/></w:tblGrid>{"".join(rows)}</w:tbl>')
        p._p.addnext(tbl)
        q = d.add_paragraph(f"after {name}")
        q.paragraph_format.space_after = Pt(0)
    d.save(path)


def exp_trheight(path):
    """Does trHeight (atLeast) include cell margins and border bands?"""
    from docx import Document
    from docx.oxml import parse_xml
    from docx.shared import Pt
    d = Document()
    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    cases = [("plain", "", ""), ("mar", '<w:top w:w="400" w:type="dxa"/><w:bottom w:w="200" w:type="dxa"/>', ""),
             ("border", "", '<w:top w:val="single" w:sz="24" w:space="0" w:color="000000"/>'),
             ("both", '<w:top w:w="400" w:type="dxa"/><w:bottom w:w="200" w:type="dxa"/>',
              '<w:top w:val="single" w:sz="24" w:space="0" w:color="000000"/>')]
    for name, mar, bd in cases:
        for rule, h in (("atLeast", 1200), ("exact", 1200)):
            p = d.add_paragraph(f"before {name} {rule}")
            p.paragraph_format.space_after = Pt(0)
            tbl = parse_xml(
                f'<w:tbl {W}><w:tblPr><w:tblW w:w="0" w:type="auto"/><w:tblCellMar>{mar}</w:tblCellMar></w:tblPr>'
                f'<w:tblGrid><w:gridCol w:w="4000"/></w:tblGrid>'
                f'<w:tr><w:trPr><w:trHeight w:val="{h}" w:hRule="{rule}"/></w:trPr>'
                f'<w:tc><w:tcPr><w:tcW w:w="4000" w:type="dxa"/><w:tcBorders>{bd}</w:tcBorders></w:tcPr>'
                f'<w:p><w:pPr><w:spacing w:after="0" w:line="280" w:lineRule="exact"/></w:pPr>'
                f'<w:r><w:t>{name} {rule} cell</w:t></w:r></w:p></w:tc></w:tr></w:tbl>')
            p._p.addnext(tbl)
            q = d.add_paragraph(f"after {name} {rule}")
            q.paragraph_format.space_after = Pt(0)
    d.save(path)


def exp_dispmath(path):
    """Line box of display equations (m:oMathPara) vs line spacing."""
    from docx import Document
    from docx.oxml import parse_xml
    from docx.shared import Pt
    M = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" ' \
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    d = Document()
    for label, ls, eq in [("single-x", 1.0, "x"), ("m115-x", 1.15, "x"), ("single-paren", 1.0, "(a)"),
                          ("m115-sub", 1.15, "S")]:
        p = d.add_paragraph(f"before {label}")
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = ls
        q = d.add_paragraph()
        q.paragraph_format.space_after = Pt(0)
        q.paragraph_format.line_spacing = ls
        if label.endswith("sub"):
            inner = ('<m:sSub><m:e><m:r><m:t>S</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub></m:sSub>')
        else:
            inner = f'<m:r><m:t>{eq}</m:t></m:r>'
        q._p.append(parse_xml(f'<m:oMathPara {M}><m:oMath>{inner}</m:oMath></m:oMathPara>'))
        r = d.add_paragraph(f"after {label}")
        r.paragraph_format.space_after = Pt(0)
        r.paragraph_format.line_spacing = ls
    d.save(path)


EXPERIMENTS = {k[4:]: v for k, v in globals().items() if k.startswith("exp_")}

if __name__ == "__main__":
    name = sys.argv[1]
    run(name, EXPERIMENTS[name])
