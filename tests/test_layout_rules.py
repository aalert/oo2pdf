"""End-to-end checks of Word layout rules found while comparing with real Word
output (sample-files.com documents, calibre's demo.docx). Each test builds a
small document with python-docx, converts it and measures the PDF.
"""
import io

import pymupdf
import pytest

import oo2pdf

docx = pytest.importorskip("docx")
from docx.oxml import parse_xml  # noqa: E402
from docx.shared import Pt  # noqa: E402

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
BORDERS = ('<w:tblBorders>' + "".join(
    f'<w:{s} w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
    for s in ("top", "left", "bottom", "right", "insideH", "insideV")) + '</w:tblBorders>')


def _convert(build):
    d = docx.Document()
    build(d)
    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)
    return pymupdf.open(stream=oo2pdf.convert(buf, kind="docx"), filetype="pdf")


def _para(text, extra=""):
    return (f'<w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>{extra}'
            f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:p>')


def _table(rows_xml, ncols=2, width=4000, tblpr=""):
    grid = "".join(f'<w:gridCol w:w="{width}"/>' for _ in range(ncols))
    return parse_xml(f'<w:tbl {W}><w:tblPr><w:tblW w:w="0" w:type="auto"/>{BORDERS}{tblpr}</w:tblPr>'
                     f'<w:tblGrid>{grid}</w:tblGrid>{rows_xml}</w:tbl>')


def _cell(content, width=4000, tcpr=""):
    return f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/>{tcpr}</w:tcPr>{content}</w:tc>'


def _fill_page(d, lines):
    for i in range(lines):
        p = d.add_paragraph(f"filler {i}")
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
    return d.paragraphs[-1]


def _words(doc, page, word):
    return [w for w in doc[page].get_text("words") if w[4] == word]


# ---------------------------------------------------------------- tables
@pytest.mark.parametrize("height, splits", [(300, True), (1079, False)])
def test_row_with_minimum_height_does_not_split_when_room_is_short(height, splits):
    """Word moves a row whose minimum height (trHeight) does not fit in the room
    left on the page instead of splitting it."""
    def build(d):
        last = _fill_page(d, 46)
        lines = "".join(_para(f"line {i}") for i in range(5))
        row = f'<w:tr><w:trPr><w:trHeight w:val="{height}"/></w:trPr>{_cell(lines)}{_cell(_para("right"))}</w:tr>'
        last._p.addnext(_table(row))
    doc = _convert(build)
    on_first = len(_words(doc, 0, "line"))
    assert (on_first > 0) == splits


def test_hidden_empty_paragraph_between_tables_takes_no_room():
    """An empty paragraph whose mark is hidden is not displayed: the tables touch."""
    def build(d):
        p = d.add_paragraph("before")
        t1 = _table(f'<w:tr>{_cell(_para("A1"))}{_cell(_para("B1"))}</w:tr>')
        hidden = parse_xml(f'<w:p {W}><w:pPr><w:rPr><w:vanish/></w:rPr></w:pPr></w:p>')
        t2 = _table(f'<w:tr>{_cell(_para("A2"))}{_cell(_para("B2"))}</w:tr>')
        p._p.addnext(t1)
        t1.addnext(hidden)
        hidden.addnext(t2)
    doc = _convert(build)
    y1 = _words(doc, 0, "A1")[0][3]
    y2 = _words(doc, 0, "A2")[0][3]
    row_pitch = y2 - y1
    assert row_pitch < 20          # one row apart, no empty line in between


def test_vertically_merged_cell_keeps_its_shading():
    """The continuation cells of a vertical merge take the first cell's fill."""
    def build(d):
        restart = '<w:shd w:val="clear" w:color="auto" w:fill="FF0000"/><w:vMerge w:val="restart"/>'
        rows = (f'<w:tr>{_cell(_para("merged"), tcpr=restart)}'
                f'{_cell(_para("r1"))}</w:tr>'
                f'<w:tr>{_cell(_para(""), tcpr="<w:vMerge/>")}{_cell(_para("r2"))}</w:tr>')
        d.add_paragraph("x")._p.addnext(_table(rows))
    doc = _convert(build)
    r2 = _words(doc, 0, "r2")[0]
    red = [dr["rect"] for dr in doc[0].get_drawings()
           if dr.get("fill") and all(abs(a - b) < 0.02 for a, b in zip(dr["fill"], (1, 0, 0)))]
    # red covers the second row's height too
    assert any(r.y0 <= r2[1] and r.y1 >= r2[3] - 1 for r in red)


def test_table_breaking_across_pages_is_closed_with_a_border():
    """The last row on a page gets a bottom border when the table continues."""
    def build(d):
        last = _fill_page(d, 30)
        rows = "".join(f'<w:tr>{_cell(_para(f"row{i}"))}{_cell(_para("x"))}</w:tr>' for i in range(40))
        last._p.addnext(_table(rows))
    doc = _convert(build)
    rows_p1 = [w for w in doc[0].get_text("words") if w[4].startswith("row")]
    bottom_text = max(w[3] for w in rows_p1)
    horizontals = [it for dr in doc[0].get_drawings() for it in dr["items"]
                   if it[0] == "l" and abs(it[1].y - it[2].y) < 0.01 and it[1].y > bottom_text]
    assert horizontals, "the last row on the page must be closed by a horizontal border"


def test_autofit_column_grows_to_its_longest_word():
    """In an auto-fit table a column whose unbreakable word does not fit grows
    and the other columns shrink."""
    long_word = "schemas.openxmlformats.org/officeDocument/2006/relationships"

    def build(d):
        rows = (f'<w:tr>{_cell(_para("a"), 2880)}{_cell(_para(long_word), 2880)}'
                f'{_cell(_para("purpose of it"), 2880)}</w:tr>')
        d.add_paragraph("x")._p.addnext(_table(rows, ncols=3, width=2880))
    doc = _convert(build)
    words = [w for w in doc[0].get_text("words") if w[4].startswith("schemas")]
    assert len(words) == 1, "the long word must not be broken across lines"


# ---------------------------------------------------------------- paragraphs
def test_drop_cap_frame_pushes_the_text_aside():
    def build(d):
        cap = parse_xml(
            f'<w:p {W}><w:pPr><w:framePr w:dropCap="drop" w:lines="3" w:wrap="around" w:vAnchor="text" '
            f'w:hAnchor="text"/><w:spacing w:line="900" w:lineRule="exact"/></w:pPr>'
            f'<w:r><w:rPr><w:sz w:val="110"/></w:rPr><w:t>D</w:t></w:r></w:p>')
        body = d.add_paragraph("rop caps are used to emphasize the leading paragraph at the start of a section. "
                               * 4)
        body._p.addprevious(cap)
    doc = _convert(build)
    first = _words(doc, 0, "rop")[0]
    letter = [w for w in doc[0].get_text("words") if w[4] == "D"][0]
    assert first[0] >= letter[2] - 0.5        # the text starts right of the cap
    assert abs(first[3] - letter[3]) > 5      # ... and the cap spans several lines


def test_math_text_layer_is_searchable():
    """Equations use astral math letters; the PDF text must extract them intact."""
    M = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'

    def build(d):
        p = d.add_paragraph("value ")
        p._p.append(parse_xml(f'<m:oMath {M}><m:r><m:t>z</m:t></m:r></m:oMath>'))
    doc = _convert(build)
    assert "\U0001d467" in doc[0].get_text()      # mathematical italic small z
