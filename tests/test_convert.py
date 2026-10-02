"""End-to-end tests that do not need Microsoft Office.

They check layout invariants measured against Word/Excel during development
(see tests/tools for the fidelity tooling that compares with real Office).
"""
import io

import pymupdf
import pytest

import oo2pdf

docx = pytest.importorskip("docx")
openpyxl = pytest.importorskip("openpyxl")


def _lines(pdf_bytes, page=0):
    doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    out = []
    for b in doc[page].get_text("dict")["blocks"]:
        for l in b.get("lines", []):
            s = l["spans"][0]
            out.append((round(s["origin"][1], 2), round(s["origin"][0], 2),
                        "".join(x["text"] for x in l["spans"])))
    return doc, sorted(out)


def _docx(build):
    d = docx.Document()
    build(d)
    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)
    return buf


def test_docx_bytes_roundtrip():
    src = _docx(lambda d: d.add_paragraph("Hello world"))
    pdf = oo2pdf.convert(src, kind="docx")
    assert pdf.startswith(b"%PDF")
    doc, lines = _lines(pdf)
    assert lines[0][2] == "Hello world"
    assert doc[0].rect.width == pytest.approx(612) and doc[0].rect.height == pytest.approx(792)


def test_paragraph_spacing_collapses_like_word():
    from docx.shared import Pt

    def build(d):
        for after, before in ((10, 24), (12, 12)):
            p = d.add_paragraph("A")
            p.paragraph_format.space_after = Pt(after)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p = d.add_paragraph("B")
            p.paragraph_format.space_before = Pt(before)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
    _, lines = _lines(oo2pdf.convert(_docx(build), kind="docx"))
    ys = [y for y, _, _ in lines]
    pitch = ys[3] - ys[2]  # line pitch between two paragraphs with 12/12 spacing
    first = ys[1] - ys[0]
    # Word uses max(after, before) unless doNotUseHTMLParagraphAutoSpacing is set
    assert first - (pitch - 12) == pytest.approx(24, abs=0.3)


def test_page_break_does_not_leave_empty_line():
    from docx.enum.text import WD_BREAK

    def build(d):
        d.add_paragraph("one")
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        d.add_paragraph("two")
    pdf = oo2pdf.convert(_docx(build), kind="docx")
    doc, lines = _lines(pdf, 1)
    _, first = _lines(pdf, 0)
    assert len(doc) == 2
    assert lines[0][2] == "two" and lines[0][0] == pytest.approx(first[0][0], abs=0.01)


def test_xlsx_basic_layout():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Left"
    ws["B1"] = 1234.5
    ws["B1"].number_format = "#,##0.00"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    pdf = oo2pdf.convert(buf, kind="xlsx")
    doc, lines = _lines(pdf)
    texts = {t for _, _, t in lines}
    assert "Left" in texts and "1,234.50" in texts
    # Excel prints Calibri 11 at 92 ppem on a 600 dpi device: 11.04 pt glyphs
    size = doc[0].get_text("dict")["blocks"][0]["lines"][0]["spans"][0]["size"]
    assert size == pytest.approx(11.04, abs=0.01)


def test_xlsx_blank_pages_are_skipped():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "x"
    ws["A200"] = "y"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    doc = pymupdf.open(stream=oo2pdf.convert(buf, kind="xlsx"), filetype="pdf")
    assert len(doc) >= 2
    assert all(p.get_text().strip() for p in doc)
